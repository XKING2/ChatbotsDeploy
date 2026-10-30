from __future__ import annotations

import json
import sys
import time
from pathlib import Path
from typing import Any

import pandas as pd

# ============================================================
# KONFIGURASI
# Script disimpan di: Chatbot/ragas_eval/retrieve_dataset.py
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

INPUT_FILE = PROJECT_ROOT / "data" / "Excel" / "Dataset_Ragas.xlsx"
OUTPUT_FILE = PROJECT_ROOT / "data" / "Excel" / "ragas_dataset_retrieved.xlsx"
CHECKPOINT_FILE = PROJECT_ROOT / "data" / "Excel" / "ragas_dataset_retrieval_checkpoint.xlsx"

# Nama sheet pada file Dataset_Ragas.xlsx
INPUT_SHEETS = {
    "Indonesia": "Indonesia",
    "English": "English",
    "日本語": "日本語",
    "中文": "中文",
    "Русский": "Русский",
}

# Beri jeda kecil antarpertanyaan agar proses lebih stabil.
DELAY_SECONDS = 0.2

# Simpan checkpoint setiap N pertanyaan.
CHECKPOINT_EVERY = 5

# Jumlah context teratas yang disimpan.
# None = simpan seluruh context yang dikembalikan retrieval pipeline.
MAX_CONTEXTS = None


def load_questions() -> pd.DataFrame:
    """Baca seluruh sheet bahasa dan gabungkan menjadi satu dataset."""
    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"File dataset tidak ditemukan: {INPUT_FILE}\n"
            "Letakkan Dataset_Ragas.xlsx di folder data/Excel."
        )

    frames = []
    excel = pd.ExcelFile(INPUT_FILE)

    for language, sheet_name in INPUT_SHEETS.items():
        if sheet_name not in excel.sheet_names:
            print(f"[PERINGATAN] Sheet '{sheet_name}' tidak ditemukan, dilewati.")
            continue

        df = pd.read_excel(INPUT_FILE, sheet_name=sheet_name)
        required = {"question_id", "topic", "question"}
        missing = required - set(df.columns)
        if missing:
            raise ValueError(
                f"Kolom pada sheet '{sheet_name}' tidak lengkap. "
                f"Kolom yang hilang: {sorted(missing)}"
            )

        df = df[["question_id", "topic", "question"]].copy()
        df["language"] = language
        df["question_id"] = df["question_id"].fillna("").astype(str).str.strip()
        df["topic"] = df["topic"].fillna("").astype(str).str.strip()
        df["question"] = df["question"].fillna("").astype(str).str.strip()
        df = df[df["question"].ne("")].copy()
        frames.append(df)

    if not frames:
        raise ValueError("Tidak ada sheet pertanyaan yang berhasil dibaca.")

    result = pd.concat(frames, ignore_index=True)
    result = result[["question_id", "language", "topic", "question"]]
    result.insert(0, "row_id", range(1, len(result) + 1))
    return result


def _is_nonempty(value: Any) -> bool:
    return value is not None and str(value).strip() != ""


def _extract_context_text(item: Any) -> str | None:
    """
    Ambil teks context dari bentuk hasil retrieval yang umum digunakan.
    Mendukung string, dict dengan content/text/document/page_content, dan object.
    """
    if item is None:
        return None

    if isinstance(item, str):
        value = item.strip()
        return value or None

    if isinstance(item, dict):
        for key in ("content", "page_content", "text", "document", "chunk_text"):
            value = item.get(key)
            if isinstance(value, str) and value.strip():
                return value.strip()

        # Bentuk ChromaDB dapat menyimpan teks di field documents.
        value = item.get("documents")
        if isinstance(value, list):
            texts = [str(x).strip() for x in value if _is_nonempty(x)]
            return "\n\n".join(texts) if texts else None

        return None

    for attr in ("page_content", "content", "text", "document"):
        value = getattr(item, attr, None)
        if isinstance(value, str) and value.strip():
            return value.strip()

    return None


def extract_contexts(retrieval_result: Any) -> list[str]:
    """
    Ekstrak context dari hasil RetrievalPipeline.retrieve().
    Memprioritaskan relevant_chunks jika tersedia, lalu mencoba struktur
    hasil retrieval lain tanpa mengambil metadata entity sebagai context.
    """
    if retrieval_result is None:
        return []

    # Bila pipeline mengembalikan list context langsung.
    if isinstance(retrieval_result, list):
        candidates = retrieval_result
    elif isinstance(retrieval_result, dict):
        candidates = None
        for key in (
            "relevant_chunks",
            "retrieved_contexts",
            "contexts",
            "chunks",
            "documents",
        ):
            value = retrieval_result.get(key)
            if isinstance(value, list) and value:
                candidates = value
                break

        # Struktur bertingkat yang mungkin digunakan oleh pipeline.
        if candidates is None:
            for key in ("results", "retrieval", "data"):
                value = retrieval_result.get(key)
                if isinstance(value, dict):
                    nested = extract_contexts(value)
                    if nested:
                        return nested
                elif isinstance(value, list) and value:
                    candidates = value
                    break

        if candidates is None:
            # Hasil satu chunk tunggal.
            one = _extract_context_text(retrieval_result)
            return [one] if one else []
    else:
        candidates = getattr(retrieval_result, "relevant_chunks", None)
        if not isinstance(candidates, list):
            candidates = getattr(retrieval_result, "contexts", None)
        if not isinstance(candidates, list):
            one = _extract_context_text(retrieval_result)
            return [one] if one else []

    contexts = []
    for item in candidates:
        text = _extract_context_text(item)
        if text and text not in contexts:
            contexts.append(text)

    if MAX_CONTEXTS is not None:
        contexts = contexts[:MAX_CONTEXTS]

    return contexts


def save_excel(df: pd.DataFrame, path: Path) -> None:
    """Simpan hasil ke Excel dengan format yang mudah dibaca."""
    from openpyxl.styles import Alignment, Font, PatternFill

    path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        # Sheet utama kompatibel untuk tahap evaluasi berikutnya.
        visible_df = df.drop(columns=["_contexts_list"], errors="ignore")
        visible_df.to_excel(writer, index=False, sheet_name="Dataset Evaluasi")

        # Sheet khusus berisi context yang dipisah per baris.
        context_rows = []
        for _, row in df.iterrows():
            contexts = row.get("_contexts_list", [])
            if not isinstance(contexts, list):
                try:
                    contexts = json.loads(str(row.get("retrieved_contexts", "[]")))
                    if not isinstance(contexts, list):
                        contexts = []
                except (json.JSONDecodeError, TypeError):
                    contexts = []
            if not contexts:
                context_rows.append({
                    "question_id": row.get("question_id", ""),
                    "language": row.get("language", ""),
                    "topic": row.get("topic", ""),
                    "question": row.get("question", ""),
                    "context_rank": None,
                    "retrieved_context": "",
                })
            else:
                for rank, context in enumerate(contexts, start=1):
                    context_rows.append({
                        "question_id": row.get("question_id", ""),
                        "language": row.get("language", ""),
                        "topic": row.get("topic", ""),
                        "question": row.get("question", ""),
                        "context_rank": rank,
                        "retrieved_context": context,
                    })

        pd.DataFrame(context_rows).to_excel(
            writer, index=False, sheet_name="Context per Baris"
        )

        summary = (
            df.assign(context_count=df["_contexts_list"].apply(
                lambda x: len(x) if isinstance(x, list) else 0
            ))
            .groupby("language", dropna=False)
            .agg(
                jumlah_pertanyaan=("question_id", "count"),
                pertanyaan_dengan_context=("context_count", lambda x: int((x > 0).sum())),
                total_context=("context_count", "sum"),
                rata_rata_context=("context_count", "mean"),
            )
            .reset_index()
        )
        summary.to_excel(writer, index=False, sheet_name="Ringkasan")

        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            for column_cells in ws.columns:
                width = max(
                    (len(str(c.value)) for c in column_cells[:100] if c.value is not None),
                    default=12,
                )
                ws.column_dimensions[column_cells[0].column_letter].width = min(
                    max(width + 2, 12), 60
                )
            for row in ws.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)


def main() -> None:
    questions = load_questions()
    print("=" * 72)
    print("RETRIEVAL CONTEXT DATASET RAG - KANTOR IMIGRASI DENPASAR")
    print("=" * 72)
    print(f"Jumlah pertanyaan: {len(questions)}")
    print(f"File input       : {INPUT_FILE}")
    print(f"File output      : {OUTPUT_FILE}")

    # Import dilakukan setelah PROJECT_ROOT dimasukkan ke sys.path.
    from LLM.chatbot import ImmigrationChatbot

    # Constructor ini membuat RetrievalPipeline tanpa menjalankan LLM jawaban.
    chatbot = ImmigrationChatbot()
    if not hasattr(chatbot, "retrieval_pipeline"):
        raise AttributeError(
            "ImmigrationChatbot tidak memiliki atribut retrieval_pipeline. "
            "Periksa implementasi LLM/chatbot.py."
        )

    # Resume dari checkpoint jika ada.
    if CHECKPOINT_FILE.exists():
        try:
            saved = pd.read_excel(CHECKPOINT_FILE, sheet_name="Dataset Evaluasi")
            if "question_id" in saved.columns and "language" in saved.columns:
                completed = set(
                    zip(saved["question_id"].astype(str), saved["language"].astype(str))
                )
            else:
                saved = pd.DataFrame()
                completed = set()
        except Exception:
            saved = pd.DataFrame()
            completed = set()
    else:
        saved = pd.DataFrame()
        completed = set()

    pending = questions[
        ~questions.apply(
            lambda row: (str(row["question_id"]), str(row["language"])) in completed,
            axis=1,
        )
    ].copy()

    print(f"Sudah tersimpan  : {len(completed)}")
    print(f"Belum diproses   : {len(pending)}")
    print("-" * 72)

    results = [saved] if not saved.empty else []
    processed_now = 0

    for _, row in pending.iterrows():
        question = str(row["question"]).strip()
        question_id = str(row["question_id"])
        language = str(row["language"])

        try:
            # Retrieval saja. Tidak memanggil ask(), build_prompt(), atau generate_answer().
            retrieval_result = chatbot.retrieval_pipeline.retrieve(
                question,
                previous_entity=None,
            )

            contexts = extract_contexts(retrieval_result)

            # Simpan hasil retrieval mentah dalam bentuk JSON agar kompatibel
            # dengan parser retrieved_contexts pada script evaluasi.
            output_row = {
                "row_id": int(row["row_id"]),
                "question_id": question_id,
                "language": language,
                "topic": str(row["topic"]),
                "question": question,
                "retrieved_contexts": json.dumps(
                    contexts, ensure_ascii=False
                ),
                # Kolom ini sengaja kosong. Isi reference berdasarkan SOP
                # sebelum menjalankan evaluasi precision/recall.
                "reference": "",
                "context_count": len(contexts),
                "retrieval_found": bool(contexts),
                "retrieval_error": "",
                "_contexts_list": contexts,
            }

            print(
                f"[OK] {question_id:<6} | {language:<10} | "
                f"{len(contexts):>2} context"
            )

        except Exception as exc:
            output_row = {
                "row_id": int(row["row_id"]),
                "question_id": question_id,
                "language": language,
                "topic": str(row["topic"]),
                "question": question,
                "retrieved_contexts": "[]",
                "reference": "",
                "context_count": 0,
                "retrieval_found": False,
                "retrieval_error": f"{type(exc).__name__}: {exc}",
                "_contexts_list": [],
            }
            print(f"[ERROR] {question_id} | {language} | {exc}")

        results.append(pd.DataFrame([output_row]))
        processed_now += 1

        if processed_now % CHECKPOINT_EVERY == 0:
            combined = pd.concat(results, ignore_index=True)
            combined = combined.drop_duplicates(
                subset=["question_id", "language"], keep="last"
            )
            combined = combined.sort_values("row_id")
            save_excel(combined, CHECKPOINT_FILE)
            print(f"  -> Checkpoint: {len(combined)}/{len(questions)} pertanyaan")

        time.sleep(DELAY_SECONDS)

    if not results:
        raise RuntimeError("Tidak ada hasil retrieval yang dapat disimpan.")

    final_df = pd.concat(results, ignore_index=True)
    final_df = final_df.drop_duplicates(
        subset=["question_id", "language"], keep="last"
    )
    final_df = final_df.sort_values("row_id").reset_index(drop=True)

    save_excel(final_df, OUTPUT_FILE)
    save_excel(final_df, CHECKPOINT_FILE)

    total_contexts = int(final_df["context_count"].sum())
    found_count = int(final_df["retrieval_found"].sum())
    error_count = int(final_df["retrieval_error"].fillna("").astype(str).str.strip().ne("").sum())

    print("\n" + "=" * 72)
    print("SELESAI")
    print("=" * 72)
    print(f"Total pertanyaan       : {len(final_df)}")
    print(f"Pertanyaan ada context : {found_count}")
    print(f"Pertanyaan tanpa ctx   : {len(final_df) - found_count}")
    print(f"Total context          : {total_contexts}")
    print(f"Baris error            : {error_count}")
    print(f"File hasil             : {OUTPUT_FILE}")
    print(f"Checkpoint             : {CHECKPOINT_FILE}")
    print("\nCatatan:")
    print("- Sheet 'Dataset Evaluasi' sudah memiliki kolom retrieved_contexts.")
    print("- Kolom reference masih kosong dan perlu diisi dari SOP.")
    print("- Setelah reference terisi, file dapat dipakai oleh evaluator retrieval.")
    print("- Periksa sheet 'Context per Baris' untuk melihat setiap context terpisah.")


if __name__ == "__main__":
    main()
