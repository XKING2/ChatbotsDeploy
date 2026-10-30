"""
Batch generation dataset untuk evaluasi RAGAS
Chatbot: Kantor Imigrasi Denpasar
Model: mengikuti config.LLM_MODEL (wajib Gemini 3)
"""

from __future__ import annotations

import json
import sys
import time
from pathlib import Path
from typing import Any

import pandas as pd

# ============================================================
# PATH DAN KONFIGURASI
# Script dijalankan dari root proyek:
#     python -m ragas_eval.ragas_generation_runner
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

INPUT_FILE = PROJECT_ROOT / "data" / "Excel" / "Dataset_Ragas.xlsx"
OUTPUT_FILE = PROJECT_ROOT / "data" / "Excel" / "ragas_dataset_generation.xlsx"
CHECKPOINT_FILE = (
    PROJECT_ROOT / "data" / "Excel" / "ragas_dataset_generation_checkpoint.xlsx"
)

INPUT_SHEETS = {
    "Indonesia": "Indonesia",
    "English": "English",
    "日本語": "日本語",
    "中文": "中文",
    "Русский": "Русский",
}

# Sepuluh pertanyaan dengan bahasa yang sama diproses dalam satu request Gemini.
# Retrieval tetap dilakukan per pertanyaan; hanya generation yang dibatch.
BATCH_SIZE = 10
CHECKPOINT_EVERY = 1
DELAY_SECONDS = 0.1

# ============================================================
# DATASET
# ============================================================

def load_questions() -> pd.DataFrame:
    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"Dataset input tidak ditemukan: {INPUT_FILE}\n"
            "Pastikan Dataset_Ragas.xlsx berada di data/Excel."
        )

    excel = pd.ExcelFile(INPUT_FILE)
    frames = []

    for language, sheet_name in INPUT_SHEETS.items():
        if sheet_name not in excel.sheet_names:
            print(f"[WARNING] Sheet {sheet_name!r} tidak ditemukan, dilewati.")
            continue

        df = pd.read_excel(INPUT_FILE, sheet_name=sheet_name)
        required = {"question_id", "topic", "question"}
        missing = required - set(df.columns)
        if missing:
            raise ValueError(
                f"Kolom pada sheet {sheet_name!r} tidak lengkap: {sorted(missing)}"
            )

        df = df[["question_id", "topic", "question"]].copy()
        df["question_id"] = df["question_id"].fillna("").astype(str).str.strip()
        df["topic"] = df["topic"].fillna("").astype(str).str.strip()
        df["question"] = df["question"].fillna("").astype(str).str.strip()
        df["language"] = language
        df = df[df["question"].ne("")].copy()
        frames.append(df)

    if not frames:
        raise ValueError("Tidak ada sheet pertanyaan yang dapat diproses.")

    data = pd.concat(frames, ignore_index=True)
    data = data[["question_id", "language", "topic", "question"]]
    data.insert(0, "row_id", range(1, len(data) + 1))

    if data["question_id"].eq("").any():
        raise ValueError("Ditemukan question_id kosong pada dataset.")

    duplicated = data.duplicated(["question_id", "language"], keep=False)
    if duplicated.any():
        duplicates = data.loc[duplicated, ["question_id", "language"]]
        raise ValueError(
            "Ditemukan kombinasi question_id dan language duplikat:\n"
            f"{duplicates.to_string(index=False)}"
        )

    return data


def normalize_contexts(value: Any) -> list[str]:
    """Pastikan contexts tersimpan sebagai list teks."""
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        items = value
    else:
        return []

    output = []
    for item in items:
        if isinstance(item, str):
            text = item.strip()
        elif isinstance(item, dict):
            text = ""
            for key in ("content", "page_content", "text", "document"):
                candidate = item.get(key)
                if isinstance(candidate, str) and candidate.strip():
                    text = candidate.strip()
                    break
        else:
            text = ""

        if text and text not in output:
            output.append(text)
    return output


# ============================================================
# CHECKPOINT DAN EXCEL
# ============================================================

def load_checkpoint() -> pd.DataFrame:
    if not CHECKPOINT_FILE.exists():
        return pd.DataFrame()

    try:
        saved = pd.read_excel(CHECKPOINT_FILE, sheet_name="Dataset Generation")
        required = {"question_id", "language", "generation_status"}
        if not required.issubset(saved.columns):
            print("[WARNING] Struktur checkpoint tidak sesuai; mulai dari awal.")
            return pd.DataFrame()
        return saved
    except Exception as exc:
        print(f"[WARNING] Checkpoint tidak dapat dibaca: {exc}")
        return pd.DataFrame()


def save_excel(df: pd.DataFrame, path: Path) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)

    visible = df.drop(columns=["_contexts_list"], errors="ignore").copy()
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        visible.to_excel(writer, index=False, sheet_name="Dataset Generation")

        context_rows = []
        for _, row in df.iterrows():
            contexts = row.get("_contexts_list", [])
            if not isinstance(contexts, list):
                try:
                    contexts = json.loads(str(row.get("retrieved_contexts", "[]")))
                except (json.JSONDecodeError, TypeError):
                    contexts = []

            if not contexts:
                context_rows.append({
                    "question_id": row.get("question_id", ""),
                    "language": row.get("language", ""),
                    "question": row.get("question", ""),
                    "context_rank": None,
                    "retrieved_context": "",
                })
            else:
                for rank, context in enumerate(contexts, start=1):
                    context_rows.append({
                        "question_id": row.get("question_id", ""),
                        "language": row.get("language", ""),
                        "question": row.get("question", ""),
                        "context_rank": rank,
                        "retrieved_context": context,
                    })

        pd.DataFrame(context_rows).to_excel(
            writer, index=False, sheet_name="Context per Baris"
        )

        summary = (
            df.groupby("language", dropna=False)
            .agg(
                jumlah_pertanyaan=("question_id", "count"),
                berhasil=("generation_status", lambda x: int(x.isin(
                    ["success", "fallback_no_context"]
                ).sum())),
                gagal=("generation_status", lambda x: int(x.eq("error").sum())),
                pertanyaan_dengan_context=(
                    "context_count", lambda x: int((pd.to_numeric(x, errors="coerce").fillna(0) > 0).sum())
                ),
            )
            .reset_index()
        )
        summary.to_excel(writer, index=False, sheet_name="Ringkasan")

        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(wrap_text=True, vertical="center")

            for col_idx, cells in enumerate(ws.columns, start=1):
                header = str(ws.cell(1, col_idx).value or "")
                if header in {"question", "response", "retrieved_contexts", "retrieved_context"}:
                    width = 65
                else:
                    width = min(max(len(header) + 3, 14), 30)
                ws.column_dimensions[get_column_letter(col_idx)].width = width

            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)



def build_batch_prompt(
    tasks: list[dict[str, str]],
    language: str,
) -> str:
    """Gabungkan prompt individual menjadi satu permintaan generation."""
    task_blocks = []
    for task in tasks:
        task_blocks.append(
            f"QUESTION_ID: {task['question_id']}\n"
            f"PROMPT:\n{task['prompt']}"
        )

    joined_tasks = "\n\n" + ("\n\n" + "=" * 60 + "\n\n").join(task_blocks)

    return f"""
Anda adalah asisten layanan informasi Kantor Imigrasi Denpasar.
Di bawah ini terdapat beberapa tugas menjawab pertanyaan. Jawab SETIAP tugas
secara terpisah berdasarkan konteks yang tersedia di dalam prompt masing-masing.

ATURAN WAJIB:
1. Gunakan bahasa {language} untuk seluruh jawaban.
2. Jawab berdasarkan konteks pada prompt untuk QUESTION_ID terkait.
3. Jangan mencampur konteks atau jawaban antar-QUESTION_ID.
4. Jika informasi tidak tersedia dalam konteks, sampaikan keterbatasan tersebut
   secara wajar dan jangan mengarang informasi.
5. Kembalikan HANYA JSON valid berupa array objek.
6. Setiap objek wajib memiliki tepat dua properti: "question_id" dan "response".
7. Buat tepat satu objek untuk setiap QUESTION_ID yang diberikan.
8. Jangan gunakan Markdown code fence atau teks penjelasan di luar JSON.

FORMAT OUTPUT:
[
  {{"question_id": "ID_PERTANYAAN", "response": "Jawaban dalam bahasa yang diminta"}}
]

DAFTAR TUGAS:
{joined_tasks}
""".strip()


def parse_batch_response(
    raw_response: Any,
    expected_ids: list[str],
) -> dict[str, str]:
    """Parse JSON hasil Gemini dan cocokkan jawaban berdasarkan question_id."""
    if raw_response is None:
        raise ValueError("Respons Gemini kosong.")

    text = str(raw_response).strip()
    if not text:
        raise ValueError("Respons Gemini kosong.")

    # Bersihkan Markdown code fence jika model tetap menghasilkannya.
    if text.startswith("```"):
        lines = text.splitlines()
        if lines and lines[0].startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]
        text = "\n".join(lines).strip()

    # Ambil bagian JSON array jika ada teks tambahan di luar JSON.
    start = text.find("[")
    end = text.rfind("]")
    if start < 0 or end < start:
        raise ValueError(
            "Respons Gemini tidak berisi JSON array yang valid. "
            f"Awal respons: {text[:500]}"
        )

    try:
        payload = json.loads(text[start:end + 1])
    except json.JSONDecodeError as exc:
        raise ValueError(
            f"JSON dari Gemini tidak valid: {exc}. "
            f"Awal respons: {text[:500]}"
        ) from exc

    if not isinstance(payload, list):
        raise ValueError("Format JSON harus berupa array.")

    expected = {str(item) for item in expected_ids}
    answers: dict[str, str] = {}

    for index, item in enumerate(payload):
        if not isinstance(item, dict):
            raise ValueError(f"Item JSON ke-{index + 1} bukan objek.")

        question_id = str(item.get("question_id", "")).strip()
        response = item.get("response")

        if not question_id:
            raise ValueError(f"Item JSON ke-{index + 1} tidak memiliki question_id.")
        if question_id not in expected:
            raise ValueError(
                f"Gemini mengembalikan question_id tidak dikenal: {question_id}"
            )
        if question_id in answers:
            raise ValueError(
                f"Gemini mengembalikan question_id duplikat: {question_id}"
            )
        if not isinstance(response, str) or not response.strip():
            raise ValueError(
                f"Jawaban untuk {question_id} kosong atau bukan teks."
            )

        answers[question_id] = response.strip()

    missing = expected - set(answers)
    if missing:
        raise ValueError(
            "Jawaban tidak lengkap. question_id yang tidak ditemukan: "
            + ", ".join(sorted(missing))
        )

    return answers


# ============================================================
# GENERATION
# ============================================================

def main() -> None:
    from dotenv import load_dotenv
    load_dotenv()

    import config as cfg

    model_name = str(getattr(cfg, "LLM_MODEL", "") or "").strip()
    if not model_name:
        raise ValueError("LLM_MODEL belum dikonfigurasi di config.py.")
    if "gemini-3" not in model_name.lower():
        raise ValueError(
            f"Model saat ini adalah {model_name!r}, bukan Gemini 3. "
            "Atur LLM_MODEL di config.py ke model Gemini 3."
        )

    questions = load_questions()
    print("=" * 76)
    print("BATCH GENERATION DATASET RAGAS - KANTOR IMIGRASI DENPASAR")
    print("=" * 76)
    print(f"Model             : {model_name}")
    print(f"Jumlah pertanyaan : {len(questions)}")
    print(f"Batch size        : {BATCH_SIZE} pertanyaan / 1 request generation")
    print(f"Input             : {INPUT_FILE}")
    print(f"Output            : {OUTPUT_FILE}")
    print("API key           : menggunakan failover di LLM.generator")
    print("-" * 76)

    from LLM.chatbot import OUT_OF_DOMAIN_MESSAGE
    from LLM.context_formatter import format_context
    from LLM.prompt_builder import build_prompt
    from LLM.generator import generate_answer
    from LLM.rate_limiter import RateLimiter
    from retrival.retrieval_pipeline import RetrievalPipeline

    retrieval_pipeline = RetrievalPipeline()
    rate_limiter = RateLimiter(max_requests=5, period=60)

    saved = load_checkpoint()
    completed_keys = set()
    if not saved.empty:
        completed = saved[
            saved["generation_status"].isin(["success", "fallback_no_context"])
        ]
        completed_keys = set(
            zip(
                completed["question_id"].astype(str),
                completed["language"].astype(str),
            )
        )

    pending = questions[
        ~questions.apply(
            lambda row: (
                str(row["question_id"]), str(row["language"])
            ) in completed_keys,
            axis=1,
        )
    ].copy()

    print(f"Sudah selesai      : {len(completed_keys)}")
    print(f"Menunggu diproses  : {len(pending)}")
    print("-" * 76)

    result_frames = [saved] if not saved.empty else []
    total_batches = 0
    for _, language_group in pending.groupby("language", sort=False):
        total_batches += (len(language_group) + BATCH_SIZE - 1) // BATCH_SIZE

    batch_number = 0
    for language, language_group in pending.groupby("language", sort=False):
        language_group = language_group.reset_index(drop=True)

        for start in range(0, len(language_group), BATCH_SIZE):
            batch_number += 1
            batch = language_group.iloc[start:start + BATCH_SIZE].copy()
            print(
                f"\n[BATCH {batch_number}/{total_batches}] "
                f"Bahasa: {language} | {len(batch)} pertanyaan | "
                "target 1 request generation"
            )

            batch_rows: list[dict[str, Any]] = []
            generation_tasks: list[dict[str, str]] = []

            # Tahap 1: retrieval per pertanyaan (tidak memanggil LLM).
            for _, row in batch.iterrows():
                question_id = str(row["question_id"])
                question = str(row["question"]).strip()

                try:
                    retrieval_result = retrieval_pipeline.retrieve(
                        question,
                        previous_entity=None,
                    )
                    found = bool(retrieval_result.get("found"))
                    context_payload = format_context(retrieval_result)

                    try:
                        parsed_context = json.loads(context_payload)
                    except (TypeError, json.JSONDecodeError):
                        parsed_context = {}

                    contexts = normalize_contexts(
                        [
                            item.get("content", "")
                            for item in (retrieval_result.get("relevant_chunks") or [])
                            if isinstance(item, dict)
                        ]
                    )

                    # Gunakan konteks yang benar-benar disusun untuk prompt.
                    has_formatted_context = (
                        isinstance(parsed_context, dict)
                        and bool(parsed_context.get("found"))
                        and bool(parsed_context.get("contexts"))
                    )

                    if found and has_formatted_context:
                        individual_prompt = build_prompt(
                            question=question,
                            context=context_payload,
                            previous_question=None,
                            previous_answer=None,
                        )
                        generation_tasks.append({
                            "question_id": question_id,
                            "prompt": individual_prompt,
                        })
                        status = "pending_generation"
                        response = ""
                    else:
                        status = "fallback_no_context"
                        response = OUT_OF_DOMAIN_MESSAGE
                        contexts = []

                    batch_rows.append({
                        "row_id": int(row["row_id"]),
                        "question_id": question_id,
                        "language": language,
                        "topic": str(row["topic"]),
                        "question": question,
                        "retrieved_contexts": json.dumps(
                            contexts, ensure_ascii=False
                        ),
                        "response": response,
                        "context_count": len(contexts),
                        "retrieval_found": bool(found and has_formatted_context),
                        "retrieval_status": str(retrieval_result.get("status") or ""),
                        "generation_status": status,
                        "generation_error": "",
                        "_contexts_list": contexts,
                    })
                    print(
                        f"  [RETRIEVED] {question_id} | "
                        f"{len(contexts)} context"
                    )

                except Exception as exc:
                    batch_rows.append({
                        "row_id": int(row["row_id"]),
                        "question_id": question_id,
                        "language": language,
                        "topic": str(row["topic"]),
                        "question": question,
                        "retrieved_contexts": "[]",
                        "response": "",
                        "context_count": 0,
                        "retrieval_found": False,
                        "retrieval_status": "error",
                        "generation_status": "error",
                        "generation_error": f"{type(exc).__name__}: {exc}",
                        "_contexts_list": [],
                    })
                    print(f"  [RETRIEVAL ERROR] {question_id}: {exc}")

            # Tahap 2: satu request Gemini untuk seluruh pertanyaan dalam batch
            # yang memiliki konteks. Semua pertanyaan dalam batch ini satu bahasa.
            if generation_tasks:
                expected_ids = [task["question_id"] for task in generation_tasks]
                try:
                    batch_prompt = build_batch_prompt(generation_tasks, language)
                    print(
                        f"  [GEMINI REQUEST] Mengirim {len(generation_tasks)} "
                        "pertanyaan dalam satu request..."
                    )
                    rate_limiter.wait()
                    batch_response = generate_answer(batch_prompt)
                    answers = parse_batch_response(batch_response, expected_ids)

                    for item in batch_rows:
                        qid = str(item["question_id"])
                        if item["generation_status"] == "pending_generation":
                            item["response"] = answers[qid]
                            item["generation_status"] = "success"

                    print(
                        f"  [BATCH OK] {len(answers)} jawaban diterima "
                        "dari satu request Gemini."
                    )

                except Exception as exc:
                    error_text = f"{type(exc).__name__}: {exc}"
                    for item in batch_rows:
                        if item["generation_status"] == "pending_generation":
                            item["generation_status"] = "error"
                            item["generation_error"] = error_text
                            item["response"] = ""
                    print(f"  [BATCH GENERATION ERROR] {error_text}")

            result_frames.append(pd.DataFrame(batch_rows))

            # Simpan setelah setiap batch supaya hasil yang sudah berhasil aman.
            combined = pd.concat(result_frames, ignore_index=True)
            combined = combined.drop_duplicates(
                subset=["question_id", "language"], keep="last"
            ).sort_values("row_id")
            save_excel(combined, CHECKPOINT_FILE)
            print(f"  [CHECKPOINT] {len(combined)} baris tersimpan.")

            time.sleep(DELAY_SECONDS)

    if not result_frames:
        raise RuntimeError("Tidak ada hasil yang dapat disimpan.")

    final_df = pd.concat(result_frames, ignore_index=True)
    final_df = final_df.drop_duplicates(
        subset=["question_id", "language"], keep="last"
    ).sort_values("row_id").reset_index(drop=True)

    save_excel(final_df, OUTPUT_FILE)
    save_excel(final_df, CHECKPOINT_FILE)

    successful = int(final_df["generation_status"].eq("success").sum())
    fallback = int(final_df["generation_status"].eq("fallback_no_context").sum())
    failed = int(final_df["generation_status"].eq("error").sum())
    context_total = int(pd.to_numeric(
        final_df["context_count"], errors="coerce"
    ).fillna(0).sum())

    print("\n" + "=" * 76)
    print("BATCH GENERATION SELESAI")
    print("=" * 76)
    print(f"Total pertanyaan       : {len(final_df)}")
    print(f"Berhasil               : {successful}")
    print(f"Fallback tanpa context : {fallback}")
    print(f"Gagal                  : {failed}")
    print(f"Total retrieved ctx    : {context_total}")
    print(f"File hasil             : {OUTPUT_FILE}")
    print(f"Checkpoint             : {CHECKPOINT_FILE}")
    print("\nCatatan:")
    print("- Retrieval tetap dijalankan per pertanyaan.")
    print("- Generation: maksimal 1 request Gemini untuk setiap 10 pertanyaan per bahasa.")
    print("- Jawaban dipasangkan kembali menggunakan question_id.")
    print("- Batch generation error akan dicoba ulang saat baris belum berstatus sukses.")
    print("- Pastikan hasil response diperiksa sebelum dipakai untuk evaluasi RAGAS.")


if __name__ == "__main__":
    main()
