"""
Evaluasi Retrieval Chatbot Imigrasi - Context Precision & Context Recall
LLM-as-a-judge (Gemini), banyak pertanyaan per satu prompt, dengan pembatas kuota.

Instalasi:
    pip install -U google-genai pandas openpyxl python-dotenv

API key (.env atau environment):
    GOOGLE_API_KEY=...            (opsional: GOOGLE_API_KEY_2 untuk project lain)

Jalankan:
    python -m ragas_eval.ragas_runner

Cara kerja:
- Satu panggilan LLM menilai sampai BATCH_SIZE pertanyaan sekaligus (default 20).
  Untuk tiap pertanyaan, juri menjawab:
    * konteks mana yang berguna untuk sampai ke reference  -> Context Precision
    * kalimat reference mana yang didukung oleh konteks    -> Context Recall
- Rumus sama dengan RAGAS:
    Context Precision = rata-rata Precision@k pada posisi konteks yang relevan
    Context Recall    = kalimat reference yang didukung / total kalimat reference
- Batas default per model: 5 request/menit, 20 request/hari, 250.000 token/menit.
  Script menunggu otomatis bila batas menit tercapai. Bila batas harian habis atau
  model error/429, batch yang sama dikirim ke model backup berikutnya.
- Checkpoint disimpan setiap panggilan. Jalankan ulang untuk melanjutkan: baris yang
  sudah sukses dilewati, dan hitungan kuota harian tersimpan di ragas_rate_state.json.
"""

from __future__ import annotations

import ast
import json
import os
import re
import time
from collections import deque
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from dotenv import load_dotenv
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# =========================
# KONFIGURASI
# =========================

PROJECT_ROOT = Path(__file__).resolve().parents[1]
INPUT_FILE = PROJECT_ROOT / "data" / "Excel" / "ragas_dataset_retrieved.xlsx"
OUTPUT_FILE = PROJECT_ROOT / "data" / "Excel" / "Hasil_Evaluasi_Ragas.xlsx"
STATE_FILE = OUTPUT_FILE.with_name("ragas_rate_state.json")
SHEET_NAME = "Dataset Evaluasi"

BATCH_SIZE = 20             # pertanyaan per prompt (maks 20)
MAX_PROMPT_TOKENS = 120_000  # batch dipotong bila perkiraan token prompt melebihi ini
OUTPUT_TOKEN_RESERVE = 6_000  # cadangan token output (termasuk thinking) per panggilan
MAX_OUTPUT_TOKENS = 40_000
THINKING = True  # False = biarkan bawaan model (boros token, risiko terpotong)
MAX_ATTEMPTS_PER_ROW = 3     # berapa kali satu baris boleh dinilai ulang bila hasilnya tidak valid
MAX_WAIT_SECONDS = 90        # tunggu model utama selama ini sebelum pindah ke backup
REQUEST_TIMEOUT = 300
RESUME = True                # lanjutkan dari hasil sukses di OUTPUT_FILE
SKIP_INVALID_ROWS = True

# Batas per model (kuota Gemini dihitung per model per project).
DEFAULT_LIMITS = {"rpm": 5, "rpd": 20, "tpm": 250_000}

# Urutan prioritas. Isi "key_env" dengan API key project lain agar kuota terpisah.
EVALUATORS = [
    {"model": "gemini-2.5-flash",       "key_env": "GOOGLE_API_KEY"},   # Utama
    {"model": "gemini-3-flash-preview", "key_env": "GOOGLE_API_KEY"},   # Backup 1
    {"model": "gemini-3.5-flash",       "key_env": "GOOGLE_API_KEY"},   # Backup 2
    {"model": "gemini-3.6-flash",       "key_env": "GOOGLE_API_KEY"},   # Backup 3
    {"model": "gemini-3.7-flash",       "key_env": "GOOGLE_API_KEY"},   # Backup 4
]

# Dapat ditimpa saat pengujian.
_now = time.time
_sleep = time.sleep


# =========================
# UTILITAS DATASET
# =========================

def clean_text(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return str(value).strip()


def parse_contexts(value: Any) -> list[str]:
    """Ubah retrieved_contexts (list / JSON string / teks) menjadi list[str]."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return []
    if isinstance(value, (list, tuple)):
        items = list(value)
    elif isinstance(value, str):
        raw = value.strip()
        if not raw:
            return []
        try:
            items = json.loads(raw)
        except json.JSONDecodeError:
            try:
                items = ast.literal_eval(raw)
            except (ValueError, SyntaxError):
                items = [raw]
    else:
        items = [str(value)]
    if not isinstance(items, (list, tuple)):
        items = [items]
    return [str(i).strip() for i in items if i is not None and str(i).strip()]


def ratio_to_percentage(value: Any) -> float | pd.NA:
    score = pd.to_numeric(value, errors="coerce")
    return pd.NA if pd.isna(score) else float(score) * 100.0


def load_dataset(path: str | Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    if not Path(path).exists():
        raise FileNotFoundError(f"File dataset tidak ditemukan: {path}")

    raw = pd.read_excel(path, sheet_name=SHEET_NAME, engine="openpyxl")
    required = ["question", "retrieved_contexts", "reference", "language", "topic"]
    missing = [c for c in required if c not in raw.columns]
    if missing:
        raise ValueError(f"Kolom wajib tidak ditemukan: {missing}")

    df = pd.DataFrame()
    df["row_id"] = raw["row_id"] if "row_id" in raw.columns else range(1, len(raw) + 1)
    df["question_id"] = raw["question_id"] if "question_id" in raw.columns else df["row_id"]
    df["language"] = raw["language"].apply(clean_text)
    df["topic"] = raw["topic"].apply(clean_text)
    df["user_input"] = raw["question"].apply(clean_text)
    df["retrieved_contexts"] = raw["retrieved_contexts"].apply(parse_contexts)
    df["reference"] = raw["reference"].apply(clean_text)

    bad = (
        df["user_input"].eq("")
        | df["reference"].eq("")
        | df["retrieved_contexts"].apply(len).eq(0)
    )
    invalid = df.loc[bad].copy()
    if len(invalid) and not SKIP_INVALID_ROWS:
        raise ValueError(f"{len(invalid)} baris kosong (pertanyaan/reference/context).")

    valid = df.loc[~bad].copy().reset_index(drop=True)
    if valid.empty:
        raise ValueError("Tidak ada baris valid yang dapat dievaluasi.")

    print("\nJumlah data berdasarkan bahasa:")
    print(valid.groupby("language").size().to_string())
    print(f"\nData valid: {len(valid)} | Data dilewati: {len(invalid)}")
    return valid, invalid


# =========================
# RUMUS METRIK (sama dengan RAGAS)
# =========================

def split_sentences(text: str) -> list[str]:
    """Pecah reference menjadi kalimat (mendukung . ! ? 。！？ dan baris baru)."""
    parts = re.split(r"(?<=[。！？!?])\s*|(?<=\.)\s+|\n+", text)
    sentences = [p.strip() for p in parts if p and p.strip()]
    return sentences or [text.strip()]


def context_precision(flags: list[int]) -> float:
    """Average precision: rata-rata Precision@k pada posisi yang relevan."""
    hits, total = 0, 0.0
    for k, flag in enumerate(flags, start=1):
        if flag:
            hits += 1
            total += hits / k
    return total / hits if hits else 0.0


def context_recall(attributed: set[int], n_sentences: int) -> float:
    return len(attributed) / n_sentences if n_sentences else 0.0


# =========================
# PROMPT & PARSING JURI
# =========================

JUDGE_INSTRUCTIONS = """You are a strict, impartial evaluator of a retrieval system.
You receive several independent ITEMS. Judge every item on its own, using only the
texts inside that item. The texts may be in Indonesian, English, Mandarin, Japanese or Russian.

For each ITEM:
1. relevant_contexts: the numbers of the CONTEXTS that are useful for arriving at the
   REFERENCE answer to the QUESTION (they contain information that supports it).
   Leave out contexts that are irrelevant or only loosely related.
2. attributed_sentences: the numbers of the REFERENCE sentences whose information can be
   found in or inferred from the CONTEXTS of that item (taken together).
   Leave out sentences that the CONTEXTS do not support.

Return ONLY valid JSON, no explanation, in exactly this shape:
{"results":[{"id":1,"relevant_contexts":[1,3],"attributed_sentences":[1,2]}]}
Write the JSON on a single line with no extra spaces or line breaks.
Include one entry for every ITEM id. Use empty lists when nothing qualifies.
"""


def build_judge_prompt(batch: list[dict]) -> tuple[str, list[dict]]:
    blocks, meta = [], []
    for local_id, row in enumerate(batch, start=1):
        sentences = split_sentences(row["reference"])
        contexts = row["retrieved_contexts"]
        meta.append({"n_sentences": len(sentences), "n_contexts": len(contexts)})
        sent = "\n".join(f"[{i}] {s}" for i, s in enumerate(sentences, 1))
        ctx = "\n".join(f"[{i}] {c}" for i, c in enumerate(contexts, 1))
        blocks.append(
            f"=== ITEM {local_id} ===\nQUESTION: {row['user_input']}\n"
            f"REFERENCE SENTENCES:\n{sent}\nCONTEXTS:\n{ctx}\n"
        )
    return JUDGE_INSTRUCTIONS + "\n" + "\n".join(blocks), meta


def _int_list(value: Any, upper: int) -> list[int]:
    if not isinstance(value, list):
        raise ValueError("bukan list")
    out = []
    for v in value:
        n = int(v)
        if 1 <= n <= upper and n not in out:
            out.append(n)
    return out


def parse_judge_response(text: str, meta: list[dict]) -> dict[int, dict]:
    """Kembalikan {local_id: {...skor}} hanya untuk item yang valid."""
    cleaned = re.sub(r"^```(?:json)?|```$", "", (text or "").strip(), flags=re.M).strip()
    try:
        data = json.loads(cleaned)
    except json.JSONDecodeError:
        data = None
        pairs = [("{", "}"), ("[", "]")]
        pairs.sort(key=lambda p: cleaned.find(p[0]) if p[0] in cleaned else 10**9)
        for open_c, close_c in pairs:
            start, end = cleaned.find(open_c), cleaned.rfind(close_c)
            if 0 <= start < end:
                try:
                    data = json.loads(cleaned[start:end + 1])
                    break
                except json.JSONDecodeError:
                    continue
        if data is None:  # balasan terpotong: ambil item yang utuh saja
            pattern = (r'\{\s*"id"\s*:\s*"?[^",}]*"?\s*,\s*"relevant_contexts"\s*:\s*\[[^\]]*\]'
                       r'\s*,\s*"attributed_sentences"\s*:\s*\[[^\]]*\]\s*\}')
            data = []
            for chunk in re.findall(pattern, cleaned):
                try:
                    data.append(json.loads(chunk))
                except json.JSONDecodeError:
                    continue
            if not data:
                return {}

    results = data.get("results") if isinstance(data, dict) else data
    parsed: dict[int, dict] = {}
    for entry in results if isinstance(results, list) else []:
        try:
            local_id = int(re.search(r"\d+", str(entry["id"])).group())
            info = meta[local_id - 1]
            relevant = _int_list(entry["relevant_contexts"], info["n_contexts"])
            attributed = _int_list(entry["attributed_sentences"], info["n_sentences"])
        except (KeyError, ValueError, TypeError, IndexError, AttributeError):
            continue
        if local_id < 1 or local_id in parsed:
            continue
        flags = [1 if i in relevant else 0 for i in range(1, info["n_contexts"] + 1)]
        parsed[local_id] = {
            "context_precision": context_precision(flags),
            "context_recall": context_recall(set(attributed), info["n_sentences"]),
            "relevant_contexts": json.dumps(relevant),
            "attributed_sentences": f"{len(attributed)}/{info['n_sentences']}",
        }
    return parsed


# =========================
# PEMBATAS KUOTA
# =========================

def estimate_tokens(text: str) -> int:
    non_ascii = sum(1 for c in text if ord(c) > 127)
    return int((len(text) - non_ascii) / 3.5 + non_ascii * 1.1) + 1


def _quota_day() -> str:
    """Kuota harian Gemini reset tengah malam waktu Pasifik (perkiraan UTC-8)."""
    try:
        from zoneinfo import ZoneInfo
        return datetime.now(ZoneInfo("America/Los_Angeles")).strftime("%Y-%m-%d")
    except Exception:
        return (datetime.now(timezone.utc) - timedelta(hours=8)).strftime("%Y-%m-%d")


class RateState:
    """Hitungan request harian per model, disimpan agar bertahan antar-run."""

    def __init__(self, path: Path):
        self.path, self.day, self.used, self.exhausted = path, _quota_day(), {}, set()
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            if data.get("day") == self.day:
                self.used = dict(data.get("used", {}))
                self.exhausted = set(data.get("exhausted", []))
        except (OSError, ValueError):
            pass

    def save(self) -> None:
        try:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            self.path.write_text(json.dumps({
                "day": self.day, "used": self.used, "exhausted": sorted(self.exhausted),
            }, indent=2), encoding="utf-8")
        except OSError:
            pass


class Evaluator:
    def __init__(self, cfg: dict, state: RateState):
        self.model = cfg["model"]
        self.key_env = cfg.get("key_env", "GOOGLE_API_KEY")
        limits = {**DEFAULT_LIMITS, **{k: cfg[k] for k in DEFAULT_LIMITS if k in cfg}}
        self.rpm, self.rpd, self.tpm = limits["rpm"], limits["rpd"], limits["tpm"]
        self.name = f"{self.key_env}:{self.model}"
        self.state = state
        self.window: deque[tuple[float, int]] = deque()  # (waktu, token)
        self.blocked_until = 0.0
        self.disabled = False
        self.errors = 0

    @property
    def api_key(self) -> str | None:
        return os.getenv(self.key_env) or (
            os.getenv("GEMINI_API_KEY") if self.key_env == "GOOGLE_API_KEY" else None
        )

    def day_exhausted(self) -> bool:
        return self.name in self.state.exhausted or self.state.used.get(self.name, 0) >= self.rpd

    def usable(self, tokens: int) -> bool:
        return not (self.disabled or self.day_exhausted() or tokens > self.tpm)

    def wait_needed(self, tokens: int) -> float:
        now = _now()
        while self.window and self.window[0][0] <= now - 60:
            self.window.popleft()
        wait = max(0.0, self.blocked_until - now)
        if len(self.window) >= self.rpm:
            wait = max(wait, self.window[0][0] + 60 - now)
        total = sum(t for _, t in self.window)
        if total + tokens > self.tpm:
            for stamp, used in self.window:
                total -= used
                if total + tokens <= self.tpm:
                    wait = max(wait, stamp + 60 - now)
                    break
        return wait + (1.0 if wait else 0.0)

    def record(self, tokens: int) -> None:
        self.window.append((_now(), tokens))

    def count_daily(self) -> None:
        self.state.used[self.name] = self.state.used.get(self.name, 0) + 1
        self.state.save()


def choose_evaluator(evaluators: list[Evaluator], tokens: int) -> tuple[Evaluator | None, float]:
    """Model prioritas tertinggi yang masih bisa dipakai; tunggu bila perlu, kalau tidak backup."""
    candidates = [e for e in evaluators if e.usable(tokens)]
    for e in candidates:
        wait = e.wait_needed(tokens)
        if wait <= MAX_WAIT_SECONDS:
            return e, wait
    if candidates:
        return min(((e, e.wait_needed(tokens)) for e in candidates), key=lambda x: x[1])
    return None, 0.0


_clients: dict[str, Any] = {}
LAST_CALL: dict[str, Any] = {}


def call_llm(ev: Evaluator, prompt: str) -> str:
    from google import genai
    from google.genai import types

    key = ev.api_key
    if key not in _clients:
        _clients[key] = genai.Client(
            api_key=key, http_options=types.HttpOptions(timeout=int(REQUEST_TIMEOUT * 1000))
        )
    def make_config(with_thinking: bool):
        opts = dict(
            temperature=0,
            response_mime_type="application/json",
            max_output_tokens=MAX_OUTPUT_TOKENS,
        )
        if with_thinking and THINKING:
            try:
                if ev.model.startswith("gemini-2.5"):
                    opts["thinking_config"] = types.ThinkingConfig(thinking_budget=0)
                else:
                    opts["thinking_config"] = types.ThinkingConfig(thinking_level="low")
            except Exception:
                pass
        return types.GenerateContentConfig(**opts)

    try:
        response = _clients[key].models.generate_content(
            model=ev.model, contents=prompt, config=make_config(True)
        )
    except Exception as error:
        if "thinking" not in str(error).lower():
            raise
        response = _clients[key].models.generate_content(  # model menolak pengaturan thinking
            model=ev.model, contents=prompt, config=make_config(False)
        )
    try:
        reason = response.candidates[0].finish_reason
    except Exception:
        reason = None
    LAST_CALL["finish_reason"] = str(reason)
    try:
        return response.text or ""
    except Exception:  # mis. diblokir / tanpa teks
        return ""


def handle_error(ev: Evaluator, error: Exception) -> str:
    """Tandai model sesuai jenis error. Mengembalikan keterangan singkat."""
    text = str(error)
    upper = text.upper()
    now = _now()
    ev.errors += 1
    if ev.errors >= 5:  # gagal beruntun: jangan terus mencoba model ini
        ev.disabled = True
        return "gagal 5x beruntun, model dinonaktifkan"
    if "PERDAY" in upper.replace(" ", "").replace("_", "") or "PER DAY" in upper:
        ev.state.exhausted.add(ev.name)
        ev.state.save()
        return "kuota harian habis"
    if "429" in upper or "RESOURCE_EXHAUSTED" in upper or "QUOTA" in upper:
        ev.blocked_until = now + 65
        return "limit per menit (429)"
    if any(m in upper for m in ("404", "NOT_FOUND", "401", "403", "API KEY", "PERMISSION_DENIED")):
        ev.disabled = True
        return "model/key tidak bisa dipakai"
    ev.blocked_until = now + 20
    return f"error sementara ({type(error).__name__})"


# =========================
# PENYIMPANAN HASIL
# =========================

def write_results(output_path, results: pd.DataFrame, invalid: pd.DataFrame,
                  summary: pd.DataFrame | None = None) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        results.to_excel(writer, sheet_name="RAGAS Results", index=False)
        if summary is not None:
            summary.to_excel(writer, sheet_name="RAGAS Summary", index=False)
        if not invalid.empty:
            invalid.drop(columns=["retrieved_contexts"], errors="ignore").to_excel(
                writer, sheet_name="Skipped Rows", index=False
            )
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for idx in range(1, ws.max_column + 1):
                header = str(ws.cell(1, idx).value or "")
                ws.column_dimensions[get_column_letter(idx)].width = (
                    24 if header.startswith("context_") else min(max(len(header) + 3, 14), 40)
                )
                if header.endswith("_persen"):
                    for r in range(2, ws.max_row + 1):
                        ws.cell(r, idx).number_format = '0.00"%"'
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)


def make_summary(results: pd.DataFrame) -> pd.DataFrame:
    if results.empty:
        return pd.DataFrame()
    groups = [
        ("Keseluruhan", results),
        *[(f"Bahasa: {k}", g) for k, g in results.groupby("language", dropna=False)],
        *[(f"Topik: {k}", g) for k, g in results.groupby("topic", dropna=False)],
    ]
    rows = []
    for label, g in groups:
        ok = g[g["status"].eq("success")]
        p = pd.to_numeric(ok["context_precision"], errors="coerce").dropna()
        r = pd.to_numeric(ok["context_recall"], errors="coerce").dropna()
        rows.append({
            "kelompok": label,
            "jumlah_data": len(g),
            "data_berhasil_dinilai": len(ok),
            "context_precision_persen": ratio_to_percentage(p.mean()) if len(p) else pd.NA,
            "context_recall_persen": ratio_to_percentage(r.mean()) if len(r) else pd.NA,
        })
    return pd.DataFrame(rows)


def load_previous_success(path: Path) -> pd.DataFrame:
    if not (RESUME and path.exists()):
        return pd.DataFrame()
    try:
        old = pd.read_excel(path, sheet_name="RAGAS Results", engine="openpyxl")
    except Exception:
        return pd.DataFrame()
    return old[old["status"].eq("success")].copy() if "status" in old.columns else pd.DataFrame()


# =========================
# PROSES UTAMA
# =========================

def take_batch(pending: deque) -> list[dict]:
    batch, tokens = [], 0
    while pending and len(batch) < BATCH_SIZE:
        row = pending[0]
        cost = estimate_tokens(
            row["user_input"] + row["reference"] + "".join(row["retrieved_contexts"])
        ) + 60
        if batch and tokens + cost > MAX_PROMPT_TOKENS:
            break
        batch.append(pending.popleft())
        tokens += cost
    return batch


def main() -> None:
    load_dotenv()
    if not 1 <= BATCH_SIZE <= 20:
        raise ValueError("BATCH_SIZE harus antara 1 dan 20.")

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    data, invalid = load_dataset(INPUT_FILE)

    state = RateState(STATE_FILE)
    evaluators = [Evaluator(cfg, state) for cfg in EVALUATORS]
    for ev in evaluators:
        if not ev.api_key:
            ev.disabled = True
            print(f"  (dilewati: {ev.key_env} belum diatur untuk {ev.model})")
    if all(ev.disabled for ev in evaluators):
        raise EnvironmentError("Tidak ada API key Gemini. Atur GOOGLE_API_KEY di .env.")

    previous = load_previous_success(OUTPUT_FILE)
    done_ids = set(previous["row_id"]) if not previous.empty else set()
    rows = [r for r in data.to_dict("records") if r["row_id"] not in done_ids]
    pending: deque = deque(rows)
    attempts: dict[Any, int] = {}
    frames = [previous] if not previous.empty else []
    failed_rows: dict[Any, dict] = {}
    call_no = 0

    print(f"Sudah selesai sebelumnya: {len(done_ids)} | Akan dinilai: {len(rows)}")
    print(f"Model: {' -> '.join(e.model for e in evaluators)}")
    print(f"Batas default: {DEFAULT_LIMITS} | Batch maks {BATCH_SIZE} pertanyaan/prompt\n")

    def checkpoint() -> None:
        parts = frames + ([pd.DataFrame(list(failed_rows.values()))] if failed_rows else [])
        all_results = pd.concat(parts, ignore_index=True) if parts else pd.DataFrame()
        if all_results.empty:
            return
        all_results = all_results.sort_values("row_id").reset_index(drop=True)
        write_results(OUTPUT_FILE, all_results, invalid, make_summary(all_results))

    stopped_early = False
    zero_streak = 0
    while pending:
        batch = take_batch(pending)
        prompt, meta = build_judge_prompt(batch)
        tokens = estimate_tokens(prompt) + OUTPUT_TOKEN_RESERVE

        ev, wait = choose_evaluator(evaluators, tokens)
        if ev is None:
            pending.extendleft(reversed(batch))
            stopped_early = True
            print("Semua model sudah mencapai batas/tidak bisa dipakai. "
                  "Hasil tersimpan; jalankan ulang setelah kuota reset untuk melanjutkan.")
            break
        if wait > 0:
            print(f"  Menunggu {wait:.0f} detik (batas per menit {ev.model})...")
            _sleep(wait)

        call_no += 1
        print(f"[Panggilan {call_no}] {ev.model} | {len(batch)} pertanyaan | ~{tokens:,} token")
        try:
            ev.record(tokens)
            text = call_llm(ev, prompt)
        except Exception as error:  # noqa: BLE001
            reason = handle_error(ev, error)
            print(f"  Gagal: {reason} -> {str(error).replace(chr(10), ' ')[:200]}")
            pending.extendleft(reversed(batch))  # batch yang sama dicoba model lain
            continue

        ev.errors = 0
        ev.count_daily()
        parsed = parse_judge_response(text, meta)
        if not parsed:
            zero_streak += 1
            debug = OUTPUT_FILE.with_name("ragas_raw_response_debug.txt")
            debug.write_text(
                f"model={ev.model} finish_reason={LAST_CALL.get('finish_reason')}\n\n{text}",
                encoding="utf-8",
            )
            print(f"  RESPONS TIDAK TERBACA | finish_reason={LAST_CALL.get('finish_reason')} "
                  f"| panjang={len(text)} | awal: {text[:300]!r}")
            print(f"  Respons lengkap disimpan di {debug}")
            if zero_streak >= 2:
                stopped_early = True
                pending.extendleft(reversed(batch))
                print("Dihentikan: 2 respons beruntun tidak terbaca (hemat kuota harian). "
                      "Kirim isi file debug di atas agar format dicek.")
                break
        else:
            zero_streak = 0
        ok_rows, retry = [], []
        for local_id, row in enumerate(batch, start=1):
            if local_id in parsed:
                ok_rows.append({
                    "row_id": row["row_id"], "question_id": row["question_id"],
                    "language": row["language"], "topic": row["topic"],
                    **parsed[local_id],
                    "context_precision_persen": ratio_to_percentage(parsed[local_id]["context_precision"]),
                    "context_recall_persen": ratio_to_percentage(parsed[local_id]["context_recall"]),
                    "batch": call_no, "evaluator_model": ev.model, "status": "success",
                })
                failed_rows.pop(row["row_id"], None)
            else:
                attempts[row["row_id"]] = attempts.get(row["row_id"], 0) + 1
                if attempts[row["row_id"]] < MAX_ATTEMPTS_PER_ROW:
                    retry.append(row)
                else:
                    failed_rows[row["row_id"]] = {
                        "row_id": row["row_id"], "question_id": row["question_id"],
                        "language": row["language"], "topic": row["topic"],
                        "context_precision": pd.NA, "context_recall": pd.NA,
                        "context_precision_persen": pd.NA, "context_recall_persen": pd.NA,
                        "batch": call_no, "evaluator_model": ev.model,
                        "status": "failed: respons juri tidak valid",
                    }
        if ok_rows:
            frames.append(pd.DataFrame(ok_rows))
        pending.extendleft(reversed(retry))
        print(f"  Berhasil {len(ok_rows)}/{len(batch)}"
              + (f" | diulang {len(retry)}" if retry else ""))
        checkpoint()

    checkpoint()
    final = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()
    print("\n========== SELESAI ==========" if not stopped_early else "\n========== BERHENTI (KUOTA) ==========")
    print(f"File hasil: {OUTPUT_FILE}")
    print(f"Berhasil: {len(final)} | Gagal: {len(failed_rows)} | Belum diproses: {len(pending)} "
          f"| Dilewati (data kosong): {len(invalid)}")
    if len(final):
        p = pd.to_numeric(final["context_precision"], errors="coerce").mean()
        r = pd.to_numeric(final["context_recall"], errors="coerce").mean()
        print(f"\nContext Precision : {ratio_to_percentage(p):.2f}%")
        print(f"Context Recall    : {ratio_to_percentage(r):.2f}%")
        print("\nPemakaian per model:")
        print(final["evaluator_model"].value_counts().to_string())


if __name__ == "__main__":
    main()